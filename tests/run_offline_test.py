#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""pdfconv 离线回归测试：在毒化代理环境下跑全部转换命令并校验产物。

用法（在 pdfconv 项目内）：
  python tests/run_offline_test.py
结果写入 artifacts/run-<timestamp>/ 与 artifacts/offline-test-report.md。
"""
import json
import os
import subprocess
import sys
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
CLI = ROOT / "pdfconv.py"
PY = sys.executable
ART = ROOT / "artifacts"
RUN = ART / f"run-{time.strftime('%Y%m%d-%H%M%S')}"
RUN.mkdir(parents=True, exist_ok=True)

# 毒化代理：若工具尝试联网，会连接失败（用于证明“离线”）
POISON = {
    "HTTP_PROXY": "http://127.0.0.1:9",
    "HTTPS_PROXY": "http://127.0.0.1:9",
    "ALL_PROXY": "http://127.0.0.1:9",
    "http_proxy": "http://127.0.0.1:9",
    "https_proxy": "http://127.0.0.1:9",
}
ENV = os.environ.copy()
ENV.update(POISON)

RESULTS = []


def run(args, timeout=600, cwd=None):
    t0 = time.time()
    cmd = [PY, str(CLI)] + [str(a) for a in args]
    p = subprocess.run(cmd, capture_output=True, text=True, encoding="utf-8", errors="replace",
                       env=ENV, timeout=timeout, cwd=str(cwd or ROOT))
    return {"cmd": cmd[2:], "rc": p.returncode, "seconds": round(time.time() - t0, 2),
            "stdout": p.stdout[-4000:], "stderr": p.stderr[-3000:]}


def check_png(path):
    from PIL import Image
    with Image.open(path) as im:
        im.load()
        return im.size[0] > 10 and im.size[1] > 10


def check_pdf_pages(path, n):
    try:
        import pymupdf as fitz
    except ImportError:  # pragma: no cover
        import fitz
    d = fitz.open(str(path))
    try:
        return d.page_count == n
    finally:
        d.close()


def check_txt(path, needles):
    t = Path(path).read_text(encoding="utf-8", errors="ignore")
    return len(t) > 10 and any(x in t for x in needles)


def check_docx(path, needles):
    from docx import Document
    d = Document(str(path))
    text = "\n".join(p.text for p in d.paragraphs)
    for tb in d.tables:
        for row in tb.rows:
            text += "\n" + " | ".join(c.text for c in row.cells)
    return any(x in text for x in needles)


def check_xlsx(path, min_sheets=3):
    from openpyxl import load_workbook
    wb = load_workbook(str(path), read_only=True, data_only=True)
    try:
        if len(wb.sheetnames) < min_sheets:
            return False
        cells = 0
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                cells += sum(1 for v in row if v not in (None, ""))
        return cells > 0
    finally:
        wb.close()


def check_pptx(path, slides):
    from pptx import Presentation
    prs = Presentation(str(path))
    return len(prs.slides) == slides


def check_images(paths, minimum=1):
    return len([p for p in paths if Path(p).exists() and Path(p).stat().st_size > 0]) >= minimum


def add(name, r, ok, note=""):
    r = dict(r)
    r.update({"name": name, "ok": bool(ok), "note": note})
    RESULTS.append(r)
    print(f"[{'PASS' if ok else 'FAIL'}] {name} ({r['seconds']}s) {note}")
    return ok


def main():
    # 0) 生成测试样本
    sample = ROOT / "tests" / "sample-multi.pdf"
    if not sample.exists():
        r = subprocess.run([PY, str(ROOT / "tests" / "make_samples.py"), str(ROOT / "tests")],
                           capture_output=True, text=True, encoding="utf-8", errors="replace", env=ENV)
        if r.returncode != 0:
            print("生成测试样本失败:", r.stderr)
            return 2
    wps_sample = ROOT / "tests" / "wps-print-template.pdf"
    d = RUN

    # 1) doctor
    r = run(["doctor", "--json"])
    deps = {}
    try:
        deps = {x["package"]: x["version"] for x in json.loads(r["stdout"])["deps"]}
    except Exception:
        pass
    add("doctor", r, r["rc"] == 0 and "PyMuPDF" in (r["stdout"] + r["stderr"]), "")

    # 2) info
    r = run(["info", str(sample)])
    add("info", r, r["rc"] == 0, "")

    # 3) pdf2png / jpg
    r = run(["pdf2png", str(sample), "--pages", "1-2", "--dpi", "120", "-o", str(d / "png")])
    pngs = sorted((d / "png").glob("*.png"))
    ok = r["rc"] == 0 and len(pngs) >= 2 and all(check_png(p) for p in pngs)
    add("pdf2png", r, ok, f"{len(pngs)} png")
    r = run(["pdf2jpg", str(sample), "--pages", "3", "--dpi", "120", "-o", str(d / "jpg")])
    jpgs = sorted((d / "jpg").glob("*.jpg"))
    add("pdf2jpg", r, r["rc"] == 0 and len(jpgs) == 1 and jpgs[0].stat().st_size > 0, f"{len(jpgs)} jpg")

    # 4) pdf2txt
    r = run(["pdf2txt", str(sample), "-o", str(d / "sample.txt")])
    add("pdf2txt", r, r["rc"] == 0 and check_txt(d / "sample.txt", ["pdfconv", "离线", "苹果", "English"]), "")

    # 5) pdf2docx
    r = run(["pdf2docx", str(sample), "-o", str(d / "sample.docx")])
    ok = r["rc"] == 0 and (d / "sample.docx").exists() and check_docx(d / "sample.docx", ["pdfconv", "离线", "苹果"])
    add("pdf2docx", r, ok, "")

    # 6) pdf2pptx
    r = run(["pdf2pptx", str(sample), "--dpi", "100", "-o", str(d / "sample.pptx")])
    ok = r["rc"] == 0 and check_pptx(d / "sample.pptx", 3)
    add("pdf2pptx", r, ok, "3 slides expected")

    # 7) pdf2xlsx
    r = run(["pdf2xlsx", str(sample), "-o", str(d / "sample.xlsx")])
    ok = r["rc"] == 0 and check_xlsx(d / "sample.xlsx", 3)
    add("pdf2xlsx", r, ok, "")

    # 8) 构造“扫描件” PDF：先渲染成图，再 img2pdf
    r = run(["pdf2png", str(sample), "--pages", "1", "--dpi", "200", "-o", str(d / "scan-src")])
    scan_png = sorted((d / "scan-src").glob("*.png"))[0]
    r2 = run(["img2pdf", str(scan_png), "-o", str(d / "scanned.pdf")])
    add("img2pdf(2step)", r2, r2["rc"] == 0 and check_pdf_pages(d / "scanned.pdf", 1), "")

    # 9) OCR
    r = run(["ocr", str(d / "scanned.pdf"), "--dpi", "200", "-o", str(d / "scanned.ocr.txt")])
    ocr_txt = d / "scanned.ocr.txt"
    ok = r["rc"] == 0 and ocr_txt.exists() and len(ocr_txt.read_text(encoding="utf-8", errors="ignore")) > 10
    note = ""
    if ok:
        t = ocr_txt.read_text(encoding="utf-8", errors="ignore")
        note = "chars=" + str(len(t))
        ok = any(x in t for x in ["pdfconv", "离线", "转换", "测试", "English", "PDF"])
    add("ocr", r, ok, note)

    # 10) 可搜索 PDF
    r = run(["pdf2searchable", str(d / "scanned.pdf"), "--dpi", "200", "-o", str(d / "scanned-searchable.pdf")])
    ok = r["rc"] == 0 and (d / "scanned-searchable.pdf").exists() and check_pdf_pages(d / "scanned-searchable.pdf", 1)
    layer = ""
    if ok:
        try:
            import pymupdf as fitz
            doc = fitz.open(str(d / "scanned-searchable.pdf"))
            layer = f"text_chars={len(doc[0].get_text('text'))}"
            doc.close()
        except Exception as e:
            layer = f"read_error={e}"
    add("pdf2searchable", r, ok, layer)

    # 11) merge
    r = run(["merge", str(sample), str(d / "scanned.pdf"), "-o", str(d / "merged.pdf")])
    add("merge", r, r["rc"] == 0 and check_pdf_pages(d / "merged.pdf", 4), "3+1 pages")

    # 12) split
    r = run(["split", str(sample), "--every", "1", "-o", str(d / "split")])
    files = sorted((d / "split").glob("*.pdf"))
    add("split", r, r["rc"] == 0 and len(files) == 3, f"{len(files)} files")

    # 13) extract-pages
    r = run(["extract-pages", str(sample), "--pages", "1,3", "-o", str(d / "pages.pdf")])
    add("extract-pages", r, r["rc"] == 0 and check_pdf_pages(d / "pages.pdf", 2), "")

    # 14) extract-images
    r = run(["extract-images", str(sample), "--min-size", "0", "-o", str(d / "images")])
    imgs = sorted((d / "images").glob("*"))
    add("extract-images", r, r["rc"] == 0 and check_images(imgs, 1), f"{len(imgs)} images")

    # 14b) pdf2imgpdf（纯图/扫描型 PDF）
    r = run(["pdf2imgpdf", str(sample), "--dpi", "100", "-o", str(d / "sample-imgpdf.pdf")])
    ok = r["rc"] == 0 and (d / "sample-imgpdf.pdf").exists() and check_pdf_pages(d / "sample-imgpdf.pdf", 3)
    add("pdf2imgpdf", r, ok, "3 pages expected")

    # 15) WPS 自带模板 PDF（真实文件，非生成样本）
    if wps_sample.exists():
        r = run(["pdf2txt", str(wps_sample), "-o", str(d / "wps-template.txt")])
        ok = r["rc"] == 0 and (d / "wps-template.txt").exists() and (d / "wps-template.txt").stat().st_size >= 0
        add("wps-template-pdf2txt", r, ok, "")

    # 16) OFD 支持（OFD→PDF/Word/图片；PDF→OFD 回环）
    ofd_sample = ROOT / "tests" / "sample-helloworld.ofd"
    if ofd_sample.exists():
        r = run(["ofd2pdf", str(ofd_sample), "-o", str(d / "ofd-sample.pdf")])
        add("ofd2pdf", r, r["rc"] == 0 and check_pdf_pages(d / "ofd-sample.pdf", 1), "1 page expected")

        r = run(["ofd2docx", str(ofd_sample), "-o", str(d / "ofd-sample.docx")])
        ok = r["rc"] == 0 and (d / "ofd-sample.docx").exists()
        note = ""
        if ok:
            try:
                from docx import Document
                docx_doc = Document(str(d / "ofd-sample.docx"))
                chars = sum(len(p.text) for p in docx_doc.paragraphs)
                note = f"paragraphs={len(docx_doc.paragraphs)} chars={chars}"
                ok = len(docx_doc.paragraphs) > 0
            except Exception as e:
                ok = False
                note = str(e)
        add("ofd2docx", r, ok, note)

        r = run(["ofd2png", str(ofd_sample), "--pages", "1", "-o", str(d / "ofd-png")])
        ofd_pngs = sorted((d / "ofd-png").glob("*.png"))
        add("ofd2png", r, r["rc"] == 0 and check_images(ofd_pngs, 1), f"{len(ofd_pngs)} png")

        r = run(["pdf2ofd", str(sample), "-o", str(d / "roundtrip.ofd")])
        ofd_round = r["rc"] == 0 and (d / "roundtrip.ofd").exists() and (d / "roundtrip.ofd").stat().st_size > 0
        add("pdf2ofd", r, ofd_round, "PDF→OFD")
        if ofd_round:
            r = run(["ofd2pdf", str(d / "roundtrip.ofd"), "-o", str(d / "roundtrip.pdf")])
            add("ofd-roundtrip", r, r["rc"] == 0 and check_pdf_pages(d / "roundtrip.pdf", 3), "3 pages expected")
    else:
        add("ofd-support", {"cmd": ["ofd2pdf"], "rc": 1, "seconds": 0, "stdout": "", "stderr": "sample missing"},
            False, "sample-helloworld.ofd missing")

    # 汇总
    passed = sum(1 for x in RESULTS if x["ok"])
    total = len(RESULTS)
    md = ["# pdfconv 离线回归测试报告", "",
          f"- 时间: {time.strftime('%Y-%m-%d %H:%M:%S')}",
          f"- Python: {sys.version.split()[0]} ({sys.executable})",
          f"- 毒化代理: {json.dumps(POISON, ensure_ascii=False)}",
          f"- 结果: **{passed}/{total} 通过**", "",
          "| 步骤 | 结果 | 耗时(s) | 备注 |", "|---|---|---|---|"]
    for x in RESULTS:
        md.append(f"| {x['name']} | {'PASS' if x['ok'] else 'FAIL'} | {x['seconds']} | {x['note']} |")
    md += ["", "## 各步骤输出（尾部）", ""]
    for x in RESULTS:
        md += [f"### {x['name']}", "```", f"cmd: {x['cmd']}", f"rc: {x['rc']}", x["stdout"].strip()[-1500:], x["stderr"].strip()[-800:], "```"]
    (ART / "offline-test-report.md").write_text("\n".join(md), encoding="utf-8")
    (RUN / "results.json").write_text(json.dumps(RESULTS, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"\n{passed}/{total} passed. Report: {ART / 'offline-test-report.md'}")
    return 0 if passed == total else 1


if __name__ == "__main__":
    sys.exit(main())
