#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""WPS 格式桥接回归测试（需要本机安装 WPS Office + pywin32）。

测试链路：
  DOCX -> WPS(.wps) -> DOCX / PDF / TXT
  PDF  -> WPS(.wps)
  OFD  -> WPS(.wps)
  WPS  -> OFD
"""
import json
import os
import subprocess
import sys
import tempfile
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
CLI = ROOT / "pdfconv.py"
ART = ROOT / "artifacts"
PY = sys.executable
OUT = Path(tempfile.mkdtemp(prefix="wps-test-"))
RESULTS = []
MARK = "CE-WPS-桥接测试-2026"


def run(args, timeout=240):
    cmd = [PY, str(CLI)] + [str(a) for a in args]
    t0 = time.time()
    p = subprocess.run(cmd, capture_output=True, text=True, encoding="utf-8",
                       errors="replace", timeout=timeout)
    return {"cmd": args, "rc": p.returncode, "out": p.stdout, "err": p.stderr,
            "sec": round(time.time() - t0, 2)}


def add(name, r, ok, note=""):
    RESULTS.append({"name": name, "ok": bool(ok), "sec": r.get("sec") if r else None,
                    "cmd": " ".join(str(x) for x in (r or {}).get("cmd", [])),
                    "stdout": (r or {}).get("out", ""), "stderr": (r or {}).get("err", ""),
                    "note": note})
    print(f"[{'PASS' if ok else 'FAIL'}] {name} ({r.get('sec') if r else '-'}s) {note}")


def is_real_wps(path):
    """wps 真格式：OLE 容器 + WpsCustomData 流。"""
    try:
        import olefile
        with open(path, "rb") as f:
            head = f.read(8)
        if head[:4] != b"\xd0\xcf\x11\xe0":
            return False
        ole = olefile.OleFileIO(str(path))
        try:
            return any(s[-1] == "WpsCustomData" for s in ole.listdir())
        finally:
            ole.close()
    except Exception:
        return False


def main():
    ART.mkdir(exist_ok=True)
    # WPS 可用性检查
    try:
        import winreg
        winreg.CloseKey(winreg.OpenKey(winreg.HKEY_CLASSES_ROOT, "KWPS.Application"))
    except Exception as e:
        print(f"SKIP: 本机没有可用的 WPS Office COM（KWPS.Application）: {e}")
        return 0

    # 1) 准备 DOCX
    from docx import Document
    src_docx = OUT / "source.docx"
    doc = Document()
    doc.add_heading("CE转换引擎 WPS 桥接测试", level=1)
    doc.add_paragraph(f"{MARK}：这是一段用于往返转换的测试文字。")
    doc.add_paragraph("第二段内容，包含数字 123456 和符号 《》。")
    doc.save(str(src_docx))

    # 2) DOCX -> WPS
    r = run(["doc2wps", src_docx, "-o", OUT / "roundtrip.wps", "--overwrite"])
    wps1 = OUT / "roundtrip.wps"
    ok = r["rc"] == 0 and wps1.exists() and is_real_wps(wps1)
    add("doc2wps", r, ok, f"size={wps1.stat().st_size if wps1.exists() else 0} 真WPS格式={is_real_wps(wps1) if wps1.exists() else False}")

    # 3) WPS -> DOCX
    r = run(["wps2docx", wps1, "-o", OUT / "back.docx", "--overwrite"])
    back = OUT / "back.docx"
    ok = r["rc"] == 0 and back.exists()
    note = ""
    if ok:
        d2 = Document(str(back))
        text = "\n".join(p.text for p in d2.paragraphs)
        note = f"paragraphs={len(d2.paragraphs)} chars={len(text)}"
        ok = MARK in text
    add("wps2docx", r, ok, note)

    # 4) WPS -> PDF
    r = run(["wps2pdf", wps1, "-o", OUT / "from-wps.pdf", "--overwrite"])
    ok = r["rc"] == 0 and (OUT / "from-wps.pdf").exists()
    note = ""
    if ok:
        try:
            import pymupdf as fitz
            d = fitz.open(str(OUT / "from-wps.pdf"))
            note = f"pages={d.page_count}"
            ok = d.page_count >= 1 and len(d[0].get_text("text")) > 0
            d.close()
        except Exception as e:
            ok, note = False, str(e)
    add("wps2pdf", r, ok, note)

    # 5) WPS -> TXT
    r = run(["wps2txt", wps1, "-o", OUT / "from-wps.txt", "--overwrite"])
    txt = OUT / "from-wps.txt"
    ok = r["rc"] == 0 and txt.exists() and MARK in txt.read_text(encoding="utf-8", errors="ignore")
    add("wps2txt", r, ok, f"chars={len(txt.read_text(encoding='utf-8', errors='ignore')) if txt.exists() else 0}")

    # 6) PDF -> WPS
    pdf_sample = ROOT / "tests" / "sample-multi.pdf"
    r = run(["pdf2wps", pdf_sample, "-o", OUT / "from-pdf.wps", "--overwrite"])
    wps2 = OUT / "from-pdf.wps"
    ok = r["rc"] == 0 and wps2.exists() and is_real_wps(wps2)
    add("pdf2wps", r, ok, f"size={wps2.stat().st_size if wps2.exists() else 0}")

    # 7) OFD -> WPS
    ofd_sample = ROOT / "tests" / "sample-helloworld.ofd"
    if ofd_sample.exists():
        r = run(["ofd2wps", ofd_sample, "-o", OUT / "from-ofd.wps", "--overwrite"])
        wps3 = OUT / "from-ofd.wps"
        ok = r["rc"] == 0 and wps3.exists() and is_real_wps(wps3)
        add("ofd2wps", r, ok, f"size={wps3.stat().st_size if wps3.exists() else 0}")

        # 8) WPS -> OFD
        r = run(["wps2ofd", wps1, "-o", OUT / "to-ofd.ofd", "--overwrite"])
        ofd = OUT / "to-ofd.ofd"
        ok = r["rc"] == 0 and ofd.exists() and ofd.read_bytes()[:2] == b"PK"
        add("wps2ofd", r, ok, f"size={ofd.stat().st_size if ofd.exists() else 0}")
    else:
        print("SKIP ofd2wps/wps2ofd: sample-helloworld.ofd missing")

    passed = sum(1 for x in RESULTS if x["ok"])
    total = len(RESULTS)
    md = ["# WPS 格式桥接测试报告", "",
          f"- 时间: {time.strftime('%Y-%m-%d %H:%M:%S')}",
          f"- Python: {sys.version.split()[0]} ({sys.executable})",
          f"- WPS: KWPS.Application（本机 WPS Office）",
          f"- 结果: **{passed}/{total} 通过**", "",
          "| 步骤 | 结果 | 耗时(s) | 备注 |", "|---|---|---|---|"]
    for x in RESULTS:
        md.append(f"| {x['name']} | {'PASS' if x['ok'] else 'FAIL'} | {x['sec']} | {x['note']} |")
    (ART / "wps-test-report.md").write_text("\n".join(md), encoding="utf-8")
    print(f"\n{passed}/{total} passed. Report: {ART / 'wps-test-report.md'}")
    return 0 if passed == total else 1


if __name__ == "__main__":
    sys.exit(main())
